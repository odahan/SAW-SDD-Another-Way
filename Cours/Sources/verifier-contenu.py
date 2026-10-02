"""Check slide/document synchronization, page boundaries and resource provenance."""
import json
import re
import hashlib
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document
from pypdf import PdfReader
from pdf2image import convert_from_path

ROOT=Path(__file__).resolve().parent.parent
V=ROOT/'Sources/Verification'
data=json.loads((ROOT/'Sources/contenu.json').read_text(encoding='utf-8'))
slides=data['slides']
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
with ZipFile(ROOT/'SAW-3.2-Cours.pptx') as z:
    parts=sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)],key=lambda n:int(re.search(r'(\d+)\.xml',n)[1]))
    assert len(parts)==len(slides)==72
    for number,s in enumerate(slides,1):
        xml=ET.fromstring(z.read(f'ppt/slides/slide{number}.xml'))
        titles=[]
        for shape in xml.findall('.//p:sp',ns):
            name=shape.find('p:nvSpPr/p:cNvPr',ns).get('name')
            if name=='title':
                titles.append(''.join(t.text or '' for t in shape.findall('.//a:t',ns)))
        assert titles==[s['title']],(number,titles,s['title'])
        notes=ET.fromstring(z.read(f'ppt/notesSlides/notesSlide{number}.xml'))
        text=' '.join(t.text or '' for t in notes.findall('.//a:t',ns))
        assert s['speech'] in text,(number,'missing trainer script')
        assert s['source'] in text,(number,'missing source')
doc=Document(ROOT/'SAW-3.2-Script-formateur.docx')
titles=[p.text for p in doc.paragraphs if re.match(r'^Slide \d{2} · ',p.text)]
assert titles==[f"Slide {s['id']:02d} · {s['title']}" for s in slides]
assert all(s['speech'] in [p.text for p in doc.paragraphs] for s in slides)
pages=json.loads((V/'word-pages.json').read_text(encoding='utf-8-sig'))
assert pages['pageCount']==74
assert [s['page'] for s in pages['slideSections']]==list(range(3,75))
native=json.loads((V/'powerpoint-layout.json').read_text(encoding='utf-8-sig'))
assert native['slideCount']==72 and not native['textOverflow']
assert sum(s['minutes'] for s in slides)==360
assert len([s for s in slides if s.get('layout')=='exercise'])==4
for lang in ['FR','EN']:
    for variant in ['SPECIFICATION','SPECIFICATION-EXEC']:
        base=ROOT/f'Ressources/SAW/{lang}/SAW-3.2-{variant}'
        md=Path(str(base)+'.md')
        txt=Path(str(base)+'.txt')
        assert md.read_text(encoding='utf-8-sig')==txt.read_text(encoding='utf-8-sig')
        assert 'Test-History.ps1' not in md.read_text(encoding='utf-8-sig')
        assert '14 bis.4' not in md.read_text(encoding='utf-8-sig')
        assert md.read_bytes()==Path(f'E:/SAW/SAW-{lang}/SAW-3.2-{variant}.md').read_bytes()
for p in (ROOT/'Ressources/TextAid').rglob('*.md'):
    original=Path('E:/TextAid')/p.relative_to(ROOT/'Ressources/TextAid')
    assert p.read_bytes()==original.read_bytes(),p
receipt=json.loads((V/'presentation-validation.json').read_text(encoding='utf-8'))
assert hashlib.sha256((ROOT/'SAW-3.2-Cours.pptx').read_bytes()).hexdigest()==receipt['finalSha256']
summary={'slides':72,'mainSlides':64,'appendixSlides':8,'guidePages':74,'minutes':360,'workshops':4,'synchronization':'passed','nativePowerPointTextboxOverflow':0,'historicalTextAidCopies':'unchanged','referenceScriptObligation':'absent','finalPptxSha256':receipt['finalSha256']}
(V/'contenu-validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False))
image_dir=V/'Word'
image_dir.mkdir(exist_ok=True)
convert_from_path(str(V/'script-formateur.pdf'),dpi=96,output_folder=str(image_dir),fmt='png',output_file='page',poppler_path=r'C:\Users\odaha\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin',paths_only=True,thread_count=1)
print('Rendered all 74 Word pages to PNG.')
