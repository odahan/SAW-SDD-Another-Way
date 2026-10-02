"""Build the slide-synchronized trainer guide and workshop resources."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / 'Sources/contenu.json').read_text(encoding='utf-8'))
SLIDES = DATA['slides']
DOC = Document()
SEC = DOC.sections[0]
SEC.page_width, SEC.page_height = Cm(21), Cm(29.7)
SEC.top_margin, SEC.bottom_margin = Cm(1.8), Cm(1.7)
SEC.left_margin = SEC.right_margin = Cm(2)
SEC.footer_distance = Cm(0.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style = DOC.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial')
DOC.styles['Normal'].font.size = Pt(11)
DOC.styles['Normal'].paragraph_format.line_spacing = 1.1
DOC.styles['Normal'].paragraph_format.space_after = Pt(6)
DOC.styles['Title'].font.size = Pt(32)
DOC.styles['Title'].font.bold = True
DOC.styles['Heading 1'].font.size = Pt(20)
DOC.styles['Heading 1'].paragraph_format.space_after = Pt(10)
DOC.styles['Heading 2'].font.size = Pt(11)
DOC.styles['Heading 2'].font.bold = True
DOC.styles['Heading 2'].paragraph_format.space_before = Pt(9)
DOC.styles['Heading 2'].paragraph_format.space_after = Pt(4)
for border in list(DOC.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
DOC.core_properties.title = 'SAW 3.2 Script du formateur'
DOC.core_properties.subject = 'Formation pratique synchronisée avec les 72 slides du cours'
DOC.core_properties.author = 'Olivier Dahan'
DOC.core_properties.keywords = 'SAW, SDD, Pro-Spec, TextAid, formation'

def para(text, style=None):
    """Append a normal paragraph with the supplied text."""
    return DOC.add_paragraph(text, style)

def part(title, text):
    """Add a compact teaching section without breaking it from its content."""
    DOC.add_heading(title,2)
    para(text)

footer = SEC.footer.paragraphs[0]
footer.text = 'SAW 3.2 · Script formateur                                            '
footer.runs[0].font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)

DOC.add_heading('SAW 3.2 Script du formateur',0)
para('Apprendre à piloter un projet avec une IA', 'Subtitle')
para('Olivier Dahan\nCours en français · 2 octobre 2026')
DOC.add_heading('Objectif de la formation',1)
para('Conduire un projet SAW depuis le recueil des informations jusqu’à la fermeture humaine des lots, avec un suivi documentaire délégué à l’IA. TextAid fournit les exemples réels de spécification, de découvertes et de validation.')
para('Ce guide accompagne SAW-3.2-Cours.pptx : 64 slides de formation et 8 annexes. Chaque fiche reprend le numéro et le titre exact de la slide, son objectif, un texte à prononcer, les indications de démonstration et les réponses attendues lorsque nécessaire.')
DOC.add_heading('Format et prérequis',1)
para('Une journée de 6 heures pédagogiques hors pauses, ou plusieurs séances. Les durées incluent les échanges, les lectures guidées et les quatre ateliers. Le texte à prononcer sert de trame ; la durée d’une fiche ne correspond pas à sa seule lecture. Prévoir les pauses séparément.')
para('Public : développeurs, responsables techniques et personnes qui définissent ou valident un produit. Savoir lire un fichier Markdown suffit pour les exercices. Aucune connaissance préalable de WPF ni aucun IDE particulier n’est nécessaire.')
DOC.add_heading('Préparation du formateur',1)
para('Ouvrir le PowerPoint et ce guide. Rendre accessibles les dossiers Ressources/SAW, Ressources/TextAid, Ressources/Ateliers et Ressources/Modeles. Distribuer les fiches d’atelier. Les démonstrations portent sur la lecture des documents : elles n’exigent ni reconstruction de TextAid, ni modèle local, ni modification du projet réel.')
para('Les fichiers TextAid sont des copies documentaires historiques. Les exemples Pro-Spec restent dans leur forme originale. Les modèles utilisent SAW 3.2. Le lot actuel de TextAid est fermé et aucun prochain lot n’est planifié : ne pas lancer le prompt sur le projet réel pour illustrer un démarrage.')
DOC.add_page_break()
DOC.add_heading('Déroulé et utilisation du guide',1)
table = DOC.add_table(rows=1,cols=3)
table.style = 'Light Shading Accent 1'
table.rows[0].cells[0].text='Module'
table.rows[0].cells[1].text='Slides'
table.rows[0].cells[2].text='Durée'
for i,(name,minutes) in enumerate(DATA['modules'],1):
    subset=[s['id'] for s in SLIDES if s['module']==i]
    row=table.add_row().cells
    row[0].text=name
    row[1].text=f'{min(subset):02d} à {max(subset):02d}'
    row[2].text=f'{minutes} min' if minutes else 'Consultation'
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(3)
            for run in p.runs:
                run.font.size=Pt(10)
                run.font.color.rgb=RGBColor(0,0,0)
                run.font.name='Arial'
DOC.add_heading('Animation et synchronisation',2)
para('Les numéros 01 à 72 correspondent au PowerPoint et à ses notes présentateur. Une fiche commence sur une nouvelle page pour faciliter le repérage pendant l’animation. Les annexes n’ajoutent pas de durée à la journée. Le diaporama reste éditable ; si une slide est ajoutée ou déplacée, mettre à jour la source commune puis régénérer les deux livrables.')
DOC.add_heading('Repères à rappeler',2)
para('Les deux prompts sont « Lis le README et exécute le prochain lot. » et « Ferme le lot en cours. ». Le protocole guide l’IA ; les validations humaines, les écarts et les arbitrages restent explicites. Une gate active FAIL ou TO TEST empêche la fermeture. N/A exige une décision humaine au ledger et des conditions applicables.')
DOC.add_heading('Sources et exemples',2)
para('SAW 3.2, révision déclarée 3.2.0-2026-09-27, copies FR/EN corrigées le 2 octobre 2026 pour retirer la section 14 bis.4. Référence historique : Pro-Spec 3 dans TextAid. Cas principaux : LOT-003 et D-008 ; LOT-010 et D-041. Les ateliers et les exemples inventés sont signalés comme pédagogiques. Les références de fin de fiche indiquent les fichiers ou clauses à consulter.')
para('Les fichiers TextAid historiques peuvent contenir des chemins vers du code ou des preuves du projet original. Les démonstrations du cours utilisent les documents copiés ; leur présence ne signifie pas que le logiciel ou tous ses résultats historiques sont distribués ici.')

for s in SLIDES:
    DOC.add_page_break()
    # Preserve exactly the visible slide title for synchronization.
    DOC.add_heading(f"Slide {s['id']:02d} · {s['title']}",1)
    duration=f"{s['minutes']} min" if s['minutes'] else 'Annexe de référence'
    p=para(f"Module {s['module']:02d} · {DATA['modules'][s['module']-1][0]} · {duration}")
    for run in p.runs:
        run.font.size=Pt(9)
    part('Objectif pédagogique',s['objective'])
    part('Texte à prononcer',s['speech'])
    if s.get('demo'):
        part('Démonstration et animation',s['demo'])
    if s.get('question'):
        part('Question au groupe',s['question'])
    if s.get('answer'):
        part('Réponse attendue ou correction',s['answer'])
    if s['id']<len(SLIDES):
        part('Transition',f"Nous allons maintenant examiner : {SLIDES[s['id']]['title']}.")
    else:
        part('Transition','Fin des annexes. Revenir au module correspondant pour reprendre le parcours de formation.')
    p=para('Sources : '+s['source'])
    for run in p.runs:
        run.font.size=Pt(9)
DOC.save(ROOT / 'SAW-3.2-Script-formateur.docx')
print(f'Created instructor guide with {len(SLIDES)} synchronized slide sections.')

# Workshop handouts contain no invented historical facts about TextAid.
workshops = {
'ATELIER-01.md': '''# Atelier 1 — Ranger les informations

Durée : 10 minutes (2 seul, 3 en binôme, 5 de correction).

Indiquer le ou les artefacts à écrire et justifier la responsabilité de chacun :

1. Windows x64 uniquement.
2. Le collage doit viser la fenêtre capturée à l’invocation.
3. Le test Replace échoue sur le candidat examiné.
4. Olivier approuve une disposition 50/50 redimensionnable avec une ligne de statut indépendante.

Cet exercice s’inspire de TextAid LOT-003. Les corrections sont dans le script formateur, slide 18.
''',
'ATELIER-02.md': '''# Atelier 2 — Spécifier Copy

Durée : 13 minutes (8 de production, 5 de mise en commun).

Scénario pédagogique distinct du projet historique TextAid : préparer LOT-001 « Copy result ».

Besoin : après une transformation réussie, l’utilisateur veut copier le résultat complet pour le coller ailleurs, sans changer le texte de l’application source. La session doit se terminer après Copy. La transformation et Replace sont hors périmètre de cet exercice.

Livrables : objectif, périmètre, 2 à 4 exigences REQ-001-nnn, cas importants, une gate AUTO et une gate HUMAN. Définir conditions et méthodes avant d’indiquer Planned. Penser à l’absence de résultat, au résultat incomplet, à l’absence de collage automatique et à la source inchangée.

Question supplémentaire : que ferait une revue LLM de cohérence ? Elle ne remplace pas le verdict humain sur l’usage.

Correction de référence dans le script formateur, slide 24.
''',
'ATELIER-03.md': '''# Atelier 3 — Quelles validations refaire

Durée : 10 minutes (5 en binômes, 5 de correction).

Scénario pédagogique :

- Candidat A : G-001-001 AUTO PASS (tests des actions Copy et Replace), G-001-002 HUMAN PASS (parcours Copy, Replace, Cancel).
- Une correction modifie le collage et produit le candidat B.
- Le développeur estime que Copy est inchangé, sans fournir encore d’examen d’impact.
- Le lot est Ready-to-close.

Décider : quels verdicts conserver dans l’historique ? Que doit identifier l’évaluation de B ? Quelles gates reviennent à TO TEST ? Que faut-il pour maintenir un PASS ? Que faire si l’impact reste indéterminé ? Quel état doit prendre le lot ?

Correction dans le script formateur, slide 44.
''',
'ATELIER-04.md': '''# Atelier 4 — Peut-on fermer ce lot

Durée : 14 minutes (8 en groupes, 6 de correction).

Scénario pédagogique fictif : LOT-001 est Ready-to-close. L’utilisateur demande « Ferme le lot en cours. ».

## Dossier présenté

SPEC : REQ-001-001 Copy complet et source inchangée ; REQ-001-002 Replace uniquement dans la fenêtre capturée ; REQ-001-003 échec sûr avec résultat copiable. Ces exigences sont actives.

Résultat : candidat B, produit après une modification du chemin de collage. A était la version évaluée précédemment.

GATES : G-001-001 AUTO PASS sur A, pas d’examen d’impact après B. G-001-002 HUMAN PASS, avec « confirmé par Alice », sans date ni résultat identifié.

FINDINGS : F-001-001 OPEN — message d’échec incompréhensible ; destination « on verra plus tard ». Aucun lot destinataire n’existe.

CONVERGENCE provisoire : TOTAL ; REQ-001-001 et REQ-001-002 satisfaites. Aucune conclusion sur REQ-001-003. Résultat obtenu : un chemin vers un fichier mutable, sans version ni copie distinguant A et B.

## Travail demandé

1. Identifier les causes de refus.
2. Proposer les actions et écritures nécessaires en préservant les traces.
3. Distinguer les actions documentaires des décisions ou validations humaines.
4. Dire quand Ready-to-close puis Closed deviennent possibles.

Ne pas inventer date, verdict ou autorisation manquants. Correction : script formateur, slides 60 et 61.
'''
}
atelier_dir=ROOT/'Ressources/Ateliers'
atelier_dir.mkdir(parents=True,exist_ok=True)
for name,content in workshops.items():
    (atelier_dir/name).write_text(content,encoding='utf-8')

models=ROOT/'Ressources/Modeles'
models.mkdir(parents=True,exist_ok=True)
templates={
'README.md': '''# SAW project protocol

> Training template. Fill every placeholder and resolve the project intent before use.

## Purpose of this file
Describe the project workflow and point to authoritative artifacts.

## Protocol reference
This project follows S.A.W. 3.2, as described in [SAW-3.2-SPECIFICATION-EXEC.md](SAW-3.2-SPECIFICATION-EXEC.md).
Revision: `3.2.0-2026-09-27`. Distributed copy corrected on 2026-10-02 to remove section 14 bis.4.

## Mandatory bootstrap
Read README, PROJECT, RULES, all active LEDGER decisions, explicitly referenced obsolete decisions, STATUS, then the lot SPEC, FINDINGS, GATES and CONVERGENCE if present. Consult HISTORY when chronology helps recovery or audit.

## Starting a lot
The human authorizes the start. Verify dependencies and the available active slot. If the next lot is ambiguous, obtain a human choice. Record In-progress before execution.

## Working on a lot
Follow SPEC and active rules. Record findings, maintain resumption information and examine validation impact after changes. Ask for required human decisions before applying them.

## Validating a lot
Evaluate each active gate according to its type on an identifiable result. Preserve previous evaluations. Record HUMAN validators and dates. N/A requires a human decision in LEDGER.

## Closing a lot
Follow the distributed protocol sections on closure. Examine every active requirement and finish every finding disposition. Prepare convergence, obtain human acceptance, finalize approved updates, verify consistency, and write Closed in STATUS last among content/state artifacts. Record the operation in HISTORY.

## Mutation rules
Preserve replaced content and identifiers. After lot start, changes of meaning to SPEC/GATES require a ledger decision. Rule evolution after initialization requires a decision with scope and effects. Do not modify a closed convergence without the resumption procedure.

## Human-only decisions
See the protocol human responsibilities and transition table. The AI may prepare these operations and must obtain the required decisions before applying them.

## History format
Append significant documentary operations to HISTORY. Preserve existing entries and use a new reconciliation entry for a correction. Use date/time and actor when known; never invent missing historical facts. The detailed EVT format is optional. No history-control script is required by the method.
''',
'PROJECT.md': '''# Project

> Training template to complete before use.

## Purpose
[Describe the product purpose.]
## Problem
[Describe the problem to solve.]
## Users and actors
[Identify users and actors.]
## Scope
[Define the intended scope.]
## Out of scope
[Define explicit exclusions.]
## Stable functional characteristics
[Describe stable characteristics.]
## Glossary
[Define project terms.]
''',
'RULES.md': '''# Project rules

> Training template. Replace the sample with actual project constraints.

## R-001 — Initial rule
Status: ACTIVE
Source: Initial project rule

Rule:
[Write the actual constraint.]

Rationale:
[Explain when useful.]

Replaces: None
Replaced by: None
''',
'STATUS.md': '''# Lot status

> Training template. Do not mark a lot Planned before its SPEC and GATES are defined.

| Lot reference | Parent | Title | Status | Replaced by | Comment |
|---|---|---|---|---|---|

## Resumption information
[Record the active lot, next action, blocking condition, expected decisions and relevant results, directly or by reference.]
''',
'LEDGER.md': '''# Durable decisions

> Training template. No decision has yet been made here. Remove the sample placeholders before use or replace them with an actual decision.

## D-001 — Decision title
Date: [Actual decision date/time]
Decided by: [Human name]
Source: [Finding or request, if relevant]
Related: [References, if relevant]

Decision:
[Actual decision]
Reason:
[Actual reason]
Consequences:
[Scope, affected lots and required actions when applicable]
Replaces: None
Replaced by: None
''',
'HISTORY.md': '''# Documentary history

> Training template. Append actual significant operations; preserve existing entries. Describe corrections in new entries. Detailed EVT identifiers are optional.

No documentary operation has been recorded in this template.
''',
'Specs/001-example/SPEC-001.md': '''# LOT-001 — Lot title

> Before working on this lot, read the repository root README.md.

## Objective
[Identifiable result]
## Context
[Project context and sources]
## In scope
[Included work]
## Out of scope
[Exclusions]
## Requirements
### REQ-001-001 — Requirement title
[Observable expected behavior]
## Important cases
[Success, failure and boundary cases]
## Dependencies
[Preconditions and lot dependencies]
## Known constraints
[Known limitations]
''',
'Specs/001-example/FINDINGS-001.md': '''# Findings — LOT-001

> Before working on this lot, read the repository root README.md.

No findings recorded yet. When a discovery is made, give it a stable F-001-nnn identifier, describe it, record its status and treatment/destinations, and preserve evidence when available. No finding may remain OPEN at closure.
''',
'Specs/001-example/GATES-001.md': '''# Gates — LOT-001

> Before working on this lot, read the repository root README.md.

## G-001-001 — Gate title
Type: [Choose AUTO, LLM or HUMAN]
Status: TO TEST
Condition:
[Closure condition]
Method:
[Evaluation procedure]
Evaluated result:
[Identify the actual examined result when evaluated]

## Test history
[Preserve prior evaluations before reevaluation. For HUMAN, record human name and actual validation date/time.]
''',
'Specs/001-example/CONVERGENCE-001.md': '''# Convergence — LOT-001

> Before working on this lot, read the repository root README.md.

Draft only. This file does not attest a closure. Complete only during the closure procedure.

Closed at: [Actual closure date/time after acceptance]
Closed by: [Human name]
Decision: [Required ledger reference, or None for an ordinary TOTAL closure]
Convergence: [TOTAL or PARTIAL]

## Result obtained
[Identify the accepted result and relate it to evaluations]
## Differences from active requirements
[Cover every active requirement and reference accepted deviations]
## Essential gate results
[Summarize valid active evaluations]
## Findings disposition
[Summarize terminal treatments and existing destinations]
## Residual work
[Explicit remaining work]
## Closure decision
[Actual human acceptance]
## Historical convergences
[Preserve previous closures during an authorized resumption]
'''
}
for name,content in templates.items():
    p=models/name
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(content,encoding='utf-8')
reference=ROOT/'Ressources/SAW/FR/SAW-3.2-SPECIFICATION-EXEC.md'
(models/reference.name).write_bytes(reference.read_bytes())

readme='''# Cours SAW 3.2

## Livrables

- `SAW-3.2-Cours.pptx` : 64 slides de formation et 8 annexes ; notes présentateur incluses.
- `SAW-3.2-Script-formateur.docx` : guide en français synchronisé avec chaque slide.
- `Ressources/` : références SAW corrigées FR/EN, copies documentaires TextAid, quatre ateliers et modèles.
- `Sources/` : contenu partagé, scripts de génération, assets et vérifications.

Durée : 6 heures pédagogiques hors pauses. Les temps incluent les échanges, lectures guidées et exercices. Les annexes sont à consulter selon les besoins.

## Utilisation

Ouvrir le PowerPoint et le DOCX. Les numéros et titres des 72 slides correspondent aux fiches du script. Les quatre exercices n’ont pas de réponses dans les fiches distribuées ; le script contient leurs corrections.

Les documents TextAid sont des copies historiques de formation, pas le projet logiciel complet. Ne pas exécuter les prompts du cours dans `E:/TextAid` : aucun prochain lot n’y est planifié. Les fichiers d’exemple peuvent conserver des liens vers des résultats et du code du projet original. Les démonstrations prévues sont documentaires.

Les modèles comportent des champs à compléter et ne constituent pas un projet opérationnel prêt à démarrer. Les scénarios inventés sont signalés comme pédagogiques.

## Mise à jour

Modifier `Sources/contenu.mjs` pour préserver la synchronisation des slides, notes et script. Les builders utilisent le runtime Node/Python fourni par Codex. Ne pas modifier les copies historiques pour fabriquer des validations. La référence SAW reste normative.

Le contenu et les illustrations s’appuient sur les documents et assets fournis par Olivier Dahan. Le livre est présenté comme ressource complémentaire : https://www.amazon.fr/dp/B0HKLSFMM3.
'''
(ROOT/'README.md').write_text(readme,encoding='utf-8')
print('Created four workshops, project templates, and course README.')
